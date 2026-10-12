"""Run after deploy: python verify_live_exports.py --api URL --code ACTIVATION_CODE.
Validates actual XLSX bytes for today, yesterday and this month.
"""
import argparse, io, json, sys, urllib.request, urllib.error
from openpyxl import load_workbook

def verify(api, code, period, location=None):
    req=urllib.request.Request(api.rstrip('/')+'/generate-excel', data=json.dumps({'activation_code':code,'report_type':'collection','period':period, **({'location':location} if location else {})}).encode(), headers={'Content-Type':'application/json'})
    with urllib.request.urlopen(req, timeout=180) as response:
        data=response.read(); typ=response.headers.get('Content-Type','')
    if 'spreadsheetml' not in typ or not data[:2] == b'PK':
        raise RuntimeError('not an Excel response')
    wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
    print(period, 'sheets:', ', '.join(wb.sheetnames), 'bytes:',len(data))
    if 'Transactions' in wb.sheetnames:
        ws=wb['Transactions']; rows=ws.iter_rows(values_only=True)
        headers=None; indices={}; count=0; paid=total=selected=0.0; fallbacks=0
        for row in rows:
            if headers is None and 'Selected Amount' in row:
                headers=list(row);indices={k:headers.index(k) for k in ('PAIDAMOUNT','TOTALAMOUNT','Selected Amount','Amount Source')};continue
            if not indices or not row or len(row)<=max(indices.values()) or not isinstance(row[indices['Selected Amount']],(int,float)):
                continue
            if str(row[0]).upper()=='GRAND TOTAL': continue
            p=float(row[indices['PAIDAMOUNT']] or 0);t=float(row[indices['TOTALAMOUNT']] or 0);v=float(row[indices['Selected Amount']] or 0)
            if abs(v-(p if p else t))>0.011: raise AssertionError('row fallback mismatch')
            paid+=p;total+=t;selected+=v;count+=1;fallbacks+=int(p==0 and t!=0)
        if not indices: raise AssertionError('detailed workbook missing accounting columns')
        if count==0: raise AssertionError('no transaction detail rows')
        print('  transactions:',count,'paid:',round(paid,2),'billed:',round(total,2),'selected:',round(selected,2),'fallback_rows:',fallbacks)
    else:
        print('  standard reconciliation format; compare to PRE-DEPLOY baseline')
    wb.close()

if __name__=='__main__':
    a=argparse.ArgumentParser();a.add_argument('--api',required=True);a.add_argument('--code',required=True);a.add_argument('--location',help='Required for monthly exports on reconciliation-mode clients');args=a.parse_args();bad=0
    for per in ('today','yesterday','this_month'):
        try: verify(args.api,args.code,per,args.location)
        except Exception as exc:
            bad+=1;print('FAIL',per,type(exc).__name__,str(exc)[:400])
    sys.exit(bool(bad))
