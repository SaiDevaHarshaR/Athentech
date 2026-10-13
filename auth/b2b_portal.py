"""Sahasra white-label B2B: registered corporate phone -> SMS OTP -> org-scoped, SELECT-only reports.

All B2B identity/session tables live in the AthenTech admin database.
No user-entered organization code or phone number is ever used as an authorization scope.
The institution-specific LIS credentials come exclusively from the trusted admin registry.
"""
from __future__ import annotations

import hashlib
import hmac
import io
import json
import os
import re
import secrets
import threading
import time
from collections import defaultdict, deque
from datetime import date, datetime, timedelta
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode
from urllib.request import Request as URLRequest, urlopen
from base64 import b64encode

from fastapi import APIRouter, Depends, HTTPException, Query, Request
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field

from auth.secrets_crypto import decrypt_secret
from database.license_db import get_conn
from database.connection import get_hospital_connection

# The default deployment refuses to issue OTPs until a real provider is configured.
TWILIO_URL = "https://verify.twilio.com/v2/Services/"
OTP_WINDOW = 600
SESSION_HOURS = 8
_lock = threading.Lock()
_hits = defaultdict(deque)


def _utc():
    return datetime.utcnow()


def _dict(row):
    if row is None:
        return None
    if hasattr(row, "items"):
        return dict(row.items())
    raise RuntimeError("Admin database must provide named columns")


def _phone(value):
    """International E.164; 10 digits without country code are treated as India (+91)."""
    p = re.sub(r"[\s().-]", "", (value or ""))
    if re.fullmatch(r"[6-9]\d{9}", p):
        p = "+91" + p
    if not re.fullmatch(r"\+[1-9]\d{7,14}", p):
        raise HTTPException(400, "Enter a mobile number with country code, such as +91XXXXXXXXXX.")
    return p


def _rate(kind, ip, institution, phone, max_attempts, window=OTP_WINDOW):
    """Additional in-process limit; Twilio Verify supplies provider-side anti-abuse controls."""
    key = (kind, str(ip)[:100], int(institution), hashlib.sha256(phone.encode()).hexdigest())
    now = time.monotonic()
    with _lock:
        bucket = _hits[key]
        while bucket and now - bucket[0] >= window:
            bucket.popleft()
        if len(bucket) >= max_attempts:
            raise HTTPException(429, "Too many attempts. Please try again later.")
        bucket.append(now)


def _ensure_tables(cur):
    # Called using a standalone admin DB connection; never in a client LIS connection.
    cur.execute("""
IF OBJECT_ID('dbo.sahasra_b2b_contacts', 'U') IS NULL
CREATE TABLE dbo.sahasra_b2b_contacts (
    id INT IDENTITY(1,1) NOT NULL PRIMARY KEY,
    institution_id INT NOT NULL,
    organisation_code NVARCHAR(80) NOT NULL,
    contact_name NVARCHAR(120) NOT NULL,
    phone_e164 NVARCHAR(20) NOT NULL,
    active BIT NOT NULL DEFAULT 1,
    created_at DATETIME2 NOT NULL DEFAULT SYSUTCDATETIME(),
    created_by NVARCHAR(100) NULL,
    CONSTRAINT UQ_sahasra_b2b_phone UNIQUE (institution_id, phone_e164)
)
""")
    cur.execute("""
IF OBJECT_ID('dbo.sahasra_b2b_sessions', 'U') IS NULL
CREATE TABLE dbo.sahasra_b2b_sessions (
    id INT IDENTITY(1,1) NOT NULL PRIMARY KEY,
    token_hash CHAR(64) NOT NULL UNIQUE,
    contact_id INT NOT NULL,
    institution_id INT NOT NULL,
    expires_at DATETIME2 NOT NULL,
    revoked BIT NOT NULL DEFAULT 0,
    created_at DATETIME2 NOT NULL DEFAULT SYSUTCDATETIME()
)
""")


def initialize_b2b():
    """Create app-owned tables, fail visibly if admin database is not ready."""
    conn = get_conn()
    try:
        _ensure_tables(conn.cursor())
        conn.commit()
    except Exception:
        try: conn.rollback()
        except Exception: pass
        raise
    finally:
        conn.close()


def _admin_conn():
    conn = get_conn()
    return conn


def _institution_by_prefix(prefix):
    if not re.fullmatch(r"[A-Za-z0-9_-]{1,50}", prefix or ""):
        return None
    conn = _admin_conn()
    try:
        row = conn.execute("SELECT * FROM institutions WHERE UPPER(client_prefix) = ?",
                           (prefix.upper(),)).fetchone()
        return _dict(row)
    finally:
        conn.close()


def _institution_by_id(institution_id):
    conn = _admin_conn()
    try:
        return _dict(conn.execute("SELECT * FROM institutions WHERE id = ?", (institution_id,)).fetchone())
    finally:
        conn.close()


def _b2b_enabled(inst):
    if not inst or str(inst.get("status") or "").lower() != "active":
        return False
    from auth.whitelabel import get_config
    cfg = get_config(int(inst["id"]))
    return bool(cfg["configured"] and cfg["enabled"] and cfg["b2b_enabled"])


def _lis(inst):
    if not inst or not inst.get("db_name"):
        raise HTTPException(503, "Institution database is not configured.")
    encrypted = inst.get("db_password")
    password = decrypt_secret(encrypted) if encrypted else None
    conn = get_hospital_connection(inst["db_name"], inst.get("db_server"), inst.get("db_user"), password)
    if conn is None:
        raise HTTPException(503, "Institution database is temporarily unavailable.")
    return conn


def _org_name(inst, org_code):
    conn = _lis(inst)
    try:
        row = conn.execute("SELECT TOP 1 ORGANISATIONNAME FROM mstOrganisation WHERE ORGANISATIONCODE = ?",
                           (org_code,)).fetchone()
        return str(row[0]) if row else None
    except Exception:
        raise HTTPException(422, "Cannot verify organisation in this LIS schema. Ask support to configure its mapping.")
    finally:
        conn.close()


def _contact(inst_id, phone):
    conn = _admin_conn()
    try:
        return _dict(conn.execute("SELECT TOP 1 * FROM dbo.sahasra_b2b_contacts WHERE institution_id=? AND phone_e164=? AND active=1",
                                  (inst_id, phone)).fetchone())
    finally:
        conn.close()


def _twilio(path, form):
    account = os.environ.get("TWILIO_ACCOUNT_SID", "")
    token = os.environ.get("TWILIO_AUTH_TOKEN", "")
    service = os.environ.get("TWILIO_VERIFY_SERVICE_SID", "")
    if not re.fullmatch(r"VA[0-9a-fA-F]{32}", service) or not account or not token:
        raise HTTPException(503, "SMS verification is not yet configured. Contact the institution administrator.")
    headers = {
        "Authorization": "Basic " + b64encode((account + ":" + token).encode()).decode(),
        "Content-Type": "application/x-www-form-urlencoded",
    }
    req = URLRequest(TWILIO_URL + service + "/" + path,
                     data=urlencode(form).encode(), headers=headers, method="POST")
    try:
        with urlopen(req, timeout=12) as response:
            return json.loads(response.read(8192))
    except HTTPError as err:
        # Never show OTP provider credentials, phone or error body to the client.
        if err.code == 429:
            raise HTTPException(429, "SMS service is busy. Please try again later.")
        if err.code in (400, 404) and path == "VerificationCheck":
            return {"status": "expired"}
        raise HTTPException(502, "SMS provider could not complete verification. Contact support.")
    except (URLError, TimeoutError, ValueError, OSError):
        raise HTTPException(503, "SMS provider is temporarily unreachable.")


class ContactIn(BaseModel):
    institution_id: int = Field(gt=0)
    organisation_code: str = Field(min_length=1, max_length=80)
    contact_name: str = Field(min_length=1, max_length=120)
    phone: str = Field(min_length=8, max_length=30)


class OTPRequest(BaseModel):
    client: str = Field(min_length=1, max_length=50)
    phone: str = Field(min_length=8, max_length=30)


class OTPCheck(OTPRequest):
    code: str = Field(min_length=4, max_length=10)


def _period_dates(p):
    today = date.today()
    if p == "today": start, end = today, today + timedelta(days=1)
    elif p == "yesterday": start, end = today - timedelta(days=1), today
    elif p == "last_7_days": start, end = today - timedelta(days=6), today + timedelta(days=1)
    elif p == "this_month": start, end = today.replace(day=1), today + timedelta(days=1)
    elif p == "last_month":
        end = today.replace(day=1)
        start = (end - timedelta(days=1)).replace(day=1)
    elif re.fullmatch(r"20\d{2}-\d{2}-\d{2}", p):
        try: start = date.fromisoformat(p)
        except ValueError: raise HTTPException(400, "Invalid report date.")
        end = start + timedelta(days=1)
    elif re.fullmatch(r"20\d{2}-\d{2}", p):
        try: start = date.fromisoformat(p + "-01")
        except ValueError: raise HTTPException(400, "Invalid report month.")
        end = date(start.year + (start.month == 12), start.month % 12 + 1, 1)
    else:
        raise HTTPException(400, "Choose today, yesterday, last_7_days, this_month, last_month, YYYY-MM, or YYYY-MM-DD.")
    return start.isoformat(), end.isoformat()


def _columns(conn, table):
    rows = conn.execute("SELECT COLUMN_NAME FROM INFORMATION_SCHEMA.COLUMNS WHERE TABLE_NAME = ?", (table,)).fetchall()
    return {str(r[0]).upper() for r in rows}


def _billing_schema(conn):
    cols = _columns(conn, "trnINVLABPRI")
    needed = {"BILLNO", "BILLDATE", "ORGANISATIONCODE"}
    if not needed.issubset(cols):
        raise HTTPException(422, "This LIS does not expose the required organisation billing columns. Contact support.")
    total = next((c for c in ("TOTALCHARGES", "TOTALAMOUNT", "GRANDTOTAL") if c in cols), None)
    if not total:
        raise HTTPException(422, "No supported billed amount column was found in this LIS.")
    # Explicitly avoid interpreting PAIDAMOUNT as collection: CentroMed and some other LISes use it differently.
    status = "p.STATUS" if "STATUS" in cols else "CAST(NULL AS VARCHAR(50))"
    return total, status


def _rows_for_org(inst, organisation_code, period, max_rows=100):
    start, end = _period_dates(period)
    conn = _lis(inst)
    try:
        amt, status = _billing_schema(conn)
        sql = ("SELECT TOP " + str(int(max_rows) + 1) + " p.BILLNO, p.BILLDATE, "
               + "p.[" + amt + "] AS BILLED, " + status + " AS BILLSTATUS "
               + "FROM trnINVLABPRI p WHERE p.ORGANISATIONCODE = ? "
               + "AND p.BILLDATE >= ? AND p.BILLDATE < ? "
               + "ORDER BY p.BILLDATE DESC, p.BILLNO DESC")
        rows = conn.execute(sql, (organisation_code, start, end)).fetchall()
        return [(str(r[0]), str(r[1])[:19], float(r[2] or 0), str(r[3] or "Not specified")) for r in rows]
    finally:
        conn.close()


def _summary_for_org(inst, org_code, period):
    start, end = _period_dates(period)
    conn = _lis(inst)
    try:
        total, _ = _billing_schema(conn)
        row = conn.execute(
            "SELECT COUNT(DISTINCT p.BILLNO), SUM(CAST(ISNULL(p.[" + total + "],0) AS FLOAT)) "
            "FROM trnINVLABPRI p WHERE p.ORGANISATIONCODE = ? "
            "AND p.BILLDATE >= ? AND p.BILLDATE < ?", (org_code, start, end)).fetchone()
        summary = {"period": period, "bills": int(row[0] or 0), "billed_total": round(float(row[1] or 0), 2),
                   "amount_label": "Billed amount (not collected cash)", "from": start, "to_exclusive": end}
        if "STATUS" in _columns(conn, "trnINVLABPRI"):
            invoice_status_rows = conn.execute(
                "SELECT p.STATUS, COUNT(DISTINCT p.BILLNO) FROM trnINVLABPRI p "
                "WHERE p.ORGANISATIONCODE=? AND p.BILLDATE >= ? AND p.BILLDATE < ? "
                "GROUP BY p.STATUS", (org_code, start, end)).fetchall()
            summary["invoice_statuses"] = [
                {"status": str(r[0] or "Not specified"), "bills": int(r[1])}
                for r in invoice_status_rows
            ]
        else:
            summary["invoice_statuses"] = []
        status_cols = _columns(conn, "trnINVLABDET")
        if {"BILLNO", "TESTSTATUS", "ORGANISATIONCODE"}.issubset(status_cols):
            statuses = conn.execute(
                "SELECT d.TESTSTATUS, COUNT(*) FROM trnINVLABDET d "
                "WHERE d.ORGANISATIONCODE=? AND EXISTS (SELECT 1 FROM trnINVLABPRI p WHERE p.BILLNO=d.BILLNO "
                "AND p.ORGANISATIONCODE=? AND p.BILLDATE >= ? AND p.BILLDATE < ?) "
                "GROUP BY d.TESTSTATUS", (org_code, org_code, start, end)).fetchall()
            summary["test_statuses"] = [{"status": str(r[0] or "Unknown"), "tests": int(r[1])} for r in statuses]
        else:
            summary["test_statuses"] = []
            summary["status_note"] = "Investigation-level status cannot be safely scoped by organisation in this LIS; invoice statuses remain available."
        return summary
    finally:
        conn.close()


def _session(token, client):
    if not token or not re.fullmatch(r"[A-Za-z0-9_-]{30,200}", token):
        raise HTTPException(401, "Sign in with your mobile number to continue.")
    inst = _institution_by_prefix(client)
    if not _b2b_enabled(inst):
        raise HTTPException(403, "B2B access is disabled for this institution.")
    conn = _admin_conn()
    try:
        row = conn.execute("""SELECT TOP 1 c.* FROM dbo.sahasra_b2b_sessions s
             JOIN dbo.sahasra_b2b_contacts c ON c.id = s.contact_id
             WHERE s.token_hash=? AND s.institution_id=? AND s.revoked=0
               AND s.expires_at > SYSUTCDATETIME() AND c.active=1
               AND c.institution_id = s.institution_id""",
             (hashlib.sha256(token.encode()).hexdigest(), int(inst["id"]))).fetchone()
        contact = _dict(row)
    finally:
        conn.close()
    if not contact:
        raise HTTPException(401, "Session expired or revoked. Please verify your phone again.")
    return inst, contact


def _auth_session(request: Request, client: str):
    authorization = request.headers.get("Authorization", "")
    scheme, _, token = authorization.partition(" ")
    if scheme.lower() != "bearer":
        raise HTTPException(401, "B2B sign-in required.")
    return _session(token.strip(), client)


def build_router(require_admin) -> APIRouter:
    router = APIRouter()

    @router.get("/admin/b2b/contacts")
    def contacts(institution_id: int = Query(..., gt=0), admin: str = Depends(require_admin)):
        if not _institution_by_id(institution_id): raise HTTPException(404, "Institution not found")
        conn = _admin_conn()
        try:
            rows = conn.execute("SELECT id, organisation_code, contact_name, phone_e164, active "
                                "FROM dbo.sahasra_b2b_contacts WHERE institution_id=? ORDER BY organisation_code, contact_name",
                                (institution_id,)).fetchall()
            return {"contacts": [{"id": r[0] if not hasattr(r, "items") else r["id"],
                      "organisation_code": r[1] if not hasattr(r, "items") else r["organisation_code"],
                      "contact_name": r[2] if not hasattr(r, "items") else r["contact_name"],
                      "phone": r[3] if not hasattr(r, "items") else r["phone_e164"],
                      "active": bool(r[4] if not hasattr(r, "items") else r["active"])} for r in rows]}
        finally: conn.close()

    @router.post("/admin/b2b/contacts")
    def create_contact(body: ContactIn, admin: str = Depends(require_admin)):
        inst = _institution_by_id(body.institution_id)
        if not inst: raise HTTPException(404, "Institution not found")
        code = body.organisation_code.strip().upper()
        if not re.fullmatch(r"[A-Za-z0-9_.-]{1,80}", code):
            raise HTTPException(400, "Invalid organisation code.")
        company = _org_name(inst, code)  # Admin cannot grant an invented organisation.
        if not company: raise HTTPException(422, "Organisation code does not exist in this client's LIS database.")
        phone = _phone(body.phone)
        conn = _admin_conn()
        try:
            conn.execute("INSERT INTO dbo.sahasra_b2b_contacts "
                         "(institution_id, organisation_code, contact_name, phone_e164, created_by) "
                         "VALUES (?,?,?,?,?)", (body.institution_id, code, body.contact_name.strip(), phone, admin))
            conn.commit()
        except Exception:
            try: conn.rollback()
            except Exception: pass
            raise HTTPException(409, "Phone is already registered for this institution or contact could not be saved.")
        finally: conn.close()
        return {"status": "success", "organisation": company, "phone": phone}

    @router.post("/admin/b2b/contacts/{contact_id}/access")
    def change_access(contact_id: int, institution_id: int = Query(..., gt=0), active: bool = Query(...),
                      admin: str = Depends(require_admin)):
        conn = _admin_conn()
        try:
            cur = conn.cursor()
            cur.execute("UPDATE dbo.sahasra_b2b_contacts SET active=? WHERE id=? AND institution_id=?",
                        (int(active), contact_id, institution_id))
            if cur.rowcount != 1: raise HTTPException(404, "Contact not found")
            if not active:
                cur.execute("UPDATE dbo.sahasra_b2b_sessions SET revoked=1 WHERE contact_id=? AND institution_id=?",
                            (contact_id, institution_id))
            conn.commit()
        except Exception:
            try: conn.rollback()
            except Exception: pass
            raise
        finally: conn.close()
        return {"status": "success", "active": active}

    @router.post("/public/b2b/request-otp")
    def request_otp(body: OTPRequest, request: Request):
        inst = _institution_by_prefix(body.client)
        if not _b2b_enabled(inst): raise HTTPException(404, "This B2B portal is not available.")
        phone = _phone(body.phone)
        ip = request.client.host if request.client else "unknown"
        _rate("start-ip", ip, inst["id"], "ip", 15)
        _rate("start-phone", "global", inst["id"], phone, 3)
        contact = _contact(int(inst["id"]), phone)
        # Same response for unknown numbers: no phone / organization enumeration.
        if contact:
            response = _twilio("Verifications", {"To": phone, "Channel": "sms"})
            if response.get("status") not in ("pending", "approved"):
                raise HTTPException(502, "SMS provider did not accept the verification request.")
        return {"status": "success", "message": "If this number is registered, a verification SMS has been sent."}

    @router.post("/public/b2b/verify-otp")
    def verify_otp(body: OTPCheck, request: Request):
        inst = _institution_by_prefix(body.client)
        if not _b2b_enabled(inst): raise HTTPException(404, "This B2B portal is not available.")
        phone = _phone(body.phone)
        if not re.fullmatch(r"\d{4,10}", body.code): raise HTTPException(400, "Enter the numeric verification code.")
        ip = request.client.host if request.client else "unknown"
        _rate("verify-ip", ip, inst["id"], "ip", 30)
        _rate("verify-phone", "global", inst["id"], phone, 8)
        contact = _contact(int(inst["id"]), phone)
        if not contact: raise HTTPException(401, "Invalid or expired verification code.")
        verified = _twilio("VerificationCheck", {"To": phone, "Code": body.code})
        if not hmac.compare_digest(str(verified.get("status", "")), "approved"):
            raise HTTPException(401, "Invalid or expired verification code.")
        token = secrets.token_urlsafe(48)
        expiry = _utc() + timedelta(hours=SESSION_HOURS)
        conn = _admin_conn()
        try:
            conn.execute("INSERT INTO dbo.sahasra_b2b_sessions "
                         "(token_hash, contact_id, institution_id, expires_at) VALUES (?,?,?,?)",
                         (hashlib.sha256(token.encode()).hexdigest(), contact["id"], inst["id"], expiry))
            conn.commit()
        finally: conn.close()
        return {"status": "success", "access_token": token, "expires_in": SESSION_HOURS * 3600,
                "contact_name": contact["contact_name"], "organisation": contact["organisation_code"]}

    @router.get("/public/b2b/overview")
    def overview(request: Request, client: str = Query(..., min_length=1, max_length=50), period: str = "this_month"):
        inst, contact = _auth_session(request, client)
        return {"status": "success", "organisation": contact["organisation_code"],
                "summary": _summary_for_org(inst, contact["organisation_code"], period)}

    @router.get("/public/b2b/bills")
    def bills(request: Request, client: str = Query(..., min_length=1, max_length=50), period: str = "this_month"):
        inst, contact = _auth_session(request, client)
        rows = _rows_for_org(inst, contact["organisation_code"], period, max_rows=100)
        return {"status": "success", "period": period,
                "bills": [{"bill_no": b, "bill_date": dt, "billed_amount": amt, "status": st}
                          for b, dt, amt, st in rows[:100]], "has_more": len(rows) > 100}

    @router.get("/public/b2b/download-excel")
    def download_excel(request: Request, client: str = Query(..., min_length=1, max_length=50), period: str = "this_month"):
        inst, contact = _auth_session(request, client)
        rows = _rows_for_org(inst, contact["organisation_code"], period, max_rows=5000)
        if len(rows) > 5000:
            raise HTTPException(413, "This report exceeds 5,000 bills. Use a narrower date range.")
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        wb = Workbook()
        ws = wb.active
        ws.title = "B2B Organisation Bills"
        ws.append(["B2B Billing Report", contact["organisation_code"], period])
        ws.append(["Invoice No", "Bill Date", "Billed Amount (INR)", "Bill Status"])
        def safe(v):
            s = str(v)
            return "'" + s if s[:1] in ("=", "+", "-", "@") else s
        for b, dt, amount, status in rows:
            ws.append([safe(b), dt, amount, safe(status)])
        for cell in ws[2]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="243B53")
        for i, width in enumerate((24, 24, 25, 25), 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "A3"
        ws.auto_filter.ref = f"A2:D{max(2, ws.max_row)}"
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        filename = "b2b_bills_" + re.sub(r"[^a-zA-Z0-9_-]", "_", period) + ".xlsx"
        return StreamingResponse(output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"',
                     "Cache-Control": "private, no-store", "X-Content-Type-Options": "nosniff"})

    @router.post("/public/b2b/logout")
    def logout(request: Request, client: str = Query(..., min_length=1, max_length=50)):
        inst, contact = _auth_session(request, client)
        token = request.headers.get("Authorization", "").partition(" ")[2].strip()
        conn = _admin_conn()
        try:
            conn.execute("UPDATE dbo.sahasra_b2b_sessions SET revoked=1 WHERE token_hash=? AND institution_id=?",
                         (hashlib.sha256(token.encode()).hexdigest(), inst["id"]))
            conn.commit()
        finally: conn.close()
        return {"status": "success"}

    return router
